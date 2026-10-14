"""
HTTP API Integration tests for Real Preview endpoints.
Verifies /api/preview/config, /api/artifacts/{id}/preview, /api/artifacts/{id}/preview/raw,
and /api/file/preview with real files and FastAPI TestClient.
"""

import os
import io
import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from fastapi.testclient import TestClient
from backend.server import app, artifact_manager, preview_service
from backend.artifacts.models import ArtifactAcquisition

class RealPreviewApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.temp_dir = Path(tempfile.mkdtemp(prefix="agy_preview_api_"))

    def test_01_preview_config_endpoint(self):
        """Verify GET /api/preview/config returns valid configuration and matrix."""
        res = self.client.get("/api/preview/config")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["preview_enabled"])
        self.assertIn(".pdf", data["support_matrix"])
        self.assertIn(".docx", data["support_matrix"])
        self.assertIn(".xlsx", data["support_matrix"])
        self.assertIn(".pptx", data["support_matrix"])
        self.assertIn(".csv", data["support_matrix"])

    def test_02_artifact_docx_preview_endpoint(self):
        """Verify DOCX preview via GET /api/artifacts/{id}/preview."""
        import docx
        doc_p = self.temp_dir / "api_sample.docx"
        doc = docx.Document()
        doc.add_heading("Production Test Title", level=1)
        doc.add_paragraph("Paragraph content for testing.")
        doc.save(str(doc_p))

        art = artifact_manager.ingest_file(doc_p, "api_sample.docx", session_id="api_sess")
        
        # Test preview payload
        res = self.client.get(f"/api/artifacts/{art.id}/preview")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "READY")
        self.assertEqual(data["preview_type"], "docx")
        self.assertIn("Production Test Title", data["content"])

        # Test raw download endpoint
        raw_res = self.client.get(f"/api/artifacts/{art.id}/content")
        self.assertEqual(raw_res.status_code, 200)
        self.assertEqual(raw_res.headers.get("content-type"), art.mime_type)

    def test_03_artifact_xlsx_preview_endpoint(self):
        """Verify XLSX multi-sheet preview via GET /api/artifacts/{id}/preview."""
        import openpyxl
        xlsx_p = self.temp_dir / "api_data.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Region", "Revenue"])
        ws.append(["EMEA", 45000])
        wb.save(str(xlsx_p))

        art = artifact_manager.ingest_file(xlsx_p, "api_data.xlsx", session_id="api_sess")
        res = self.client.get(f"/api/artifacts/{art.id}/preview")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["preview_type"], "xlsx")
        self.assertIn("Sales", data["sheet_names"])
        self.assertIn("EMEA", data["content"])

    def test_04_artifact_pdf_raw_stream_endpoint(self):
        """Verify PDF sandboxed streaming via GET /api/artifacts/{id}/preview/raw."""
        import pypdf
        pdf_p = self.temp_dir / "api_doc.pdf"
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with open(pdf_p, "wb") as f:
            writer.write(f)

        art = artifact_manager.ingest_file(pdf_p, "api_doc.pdf", session_id="api_sess")
        res = self.client.get(f"/api/artifacts/{art.id}/preview/raw")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers.get("x-content-type-options"), "nosniff")
        self.assertGreater(len(res.content), 0)

    def test_05_workspace_file_preview_endpoint(self):
        """Verify GET /api/file/preview?path=... with rich renderer."""
        csv_p = self.temp_dir / "workspace_test.csv"
        csv_p.write_text("item,qty\nServer,10\nRouter,5\n", encoding="utf-8")

        res = self.client.get(f"/api/file/preview?path={str(csv_p)}")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["preview_type"], "table")
        self.assertIn("Server", data["content"])

if __name__ == "__main__":
    unittest.main()
