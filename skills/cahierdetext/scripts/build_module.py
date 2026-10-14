"""
Automated Document Builder for Cahier de Texte.
Builds a complete, perfectly paginated Word document from master template and JSON input.
"""
import sys
import os
import json
import copy
import shutil
import docx
from docx.shared import Pt
from docx.oxml import parse_xml

def build_document(master_template, target_output, json_path, header_info=None):
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    if header_info is None:
        header_info = data['header_info']
        
    sessions = data['sessions']
    page_comments = data['page_comments']
    page_taux = data['page_taux']
    total_sessions = len(sessions)
    total_pages = (total_sessions + 3) // 4
    
    print(f"Building document: {total_sessions} sessions across {total_pages} pages.")
    
    # Fresh copy of master template
    shutil.copyfile(master_template, target_output)
    doc = docx.Document(target_output)
    body = doc._body._body
    
    proto_header_tbl = copy.deepcopy(doc.tables[0]._tbl)
    proto_session_tbl = copy.deepcopy(doc.tables[1]._tbl) # 4 sessions (5 rows)
    proto_footer_tbl = copy.deepcopy(doc.tables[2]._tbl)
    
    proto_note_p = None
    for el in body:
        if el.tag.split('}')[-1] == 'p' and 'Le formateur doit joindre' in ''.join(el.itertext()):
            proto_note_p = copy.deepcopy(el)
            break
            
    clean_spacer_p = parse_xml(r'''
    <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:pPr>
        <w:spacing w:before="0" w:after="20" w:line="100" w:lineRule="exact"/>
        <w:rPr><w:sz w:val="4"/></w:rPr>
      </w:pPr>
    </w:p>
    ''')

    clean_pagebreak_p = parse_xml(r'''
    <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:pPr>
        <w:pageBreakBefore/>
      </w:pPr>
    </w:p>
    ''')

    sectPr = copy.deepcopy(body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr'))
    
    # CLEAR body
    body.clear()
    
    for p_idx in range(total_pages):
        body.append(copy.deepcopy(proto_header_tbl))
        body.append(copy.deepcopy(clean_spacer_p))
        
        s_tbl_el = copy.deepcopy(proto_session_tbl)
        start_s = p_idx * 4
        end_s = min(start_s + 4, total_sessions)
        page_s_count = end_s - start_s
        
        if page_s_count < 4:
            tr_elements = s_tbl_el.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
            for tr in tr_elements[page_s_count + 1:]:
                s_tbl_el.remove(tr)
                
        body.append(s_tbl_el)
        body.append(copy.deepcopy(clean_spacer_p))
        
        body.append(copy.deepcopy(proto_footer_tbl))
        body.append(copy.deepcopy(proto_note_p))
        
        if p_idx < total_pages - 1:
            body.append(copy.deepcopy(clean_pagebreak_p))
        else:
            body.append(copy.deepcopy(clean_spacer_p))
            body.append(sectPr)
            
    doc.save(target_output)
    doc = docx.Document(target_output)
    
    def set_cell_simple(cell, text, bold=True, sz_pt=8):
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run = p.add_run(text)
        run.bold = bold
        if sz_pt:
            run.font.size = Pt(sz_pt)

    # Set cantSplit and tight spacing on all cells
    for tbl in doc.tables:
        for row in tbl.rows:
            trPr = row._tr.get_or_add_trPr()
            if trPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit') is None:
                trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    sp = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing')
                    if sp is not None:
                        pPr.remove(sp)
                    pPr.append(parse_xml(r'<w:spacing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:before="0" w:after="0" w:line="200" w:lineRule="auto"/>'))

    # Populate pages
    for page_idx in range(total_pages):
        h_tbl = doc.tables[page_idx * 3]
        s_tbl = doc.tables[page_idx * 3 + 1]
        f_tbl = doc.tables[page_idx * 3 + 2]
        
        # 1. Header Table
        title_text = f" ETAT D’AVANCEMENT ET DE REALISATION DU PROGRAMME DE FORMATION {header_info['module_code']} - Fiche de Suivi Pédagogique (2026-2027)"
        set_cell_simple(h_tbl.rows[0].cells[0], title_text, bold=True, sz_pt=None)
        if len(h_tbl.rows[0].cells) > 1 and h_tbl.rows[0].cells[1]._tc != h_tbl.rows[0].cells[0]._tc:
            set_cell_simple(h_tbl.rows[0].cells[1], title_text, bold=True, sz_pt=None)
            
        p_r1c0 = h_tbl.rows[1].cells[0].paragraphs[0]
        for r in list(p_r1c0.runs): p_r1c0._p.remove(r._r)
        r1 = p_r1c0.add_run("Intitulé de l'unité de formation : ")
        r1.bold = True
        r2 = p_r1c0.add_run(header_info['title'])
        r2.bold = False
        
        p_r1c1 = h_tbl.rows[1].cells[1].paragraphs[0]
        for r in list(p_r1c1.runs): p_r1c1._p.remove(r._r)
        r1 = p_r1c1.add_run("Formateur : ")
        r1.bold = True
        r2 = p_r1c1.add_run(header_info['trainer'])
        r2.bold = False
        
        p_r2c0 = h_tbl.rows[2].cells[0].paragraphs[0]
        for r in list(p_r2c0.runs): p_r2c0._p.remove(r._r)
        r1 = p_r2c0.add_run("Objectif général de l'unité de formation : ")
        r1.bold = True
        r2 = p_r2c0.add_run(header_info['general_obj'])
        r2.bold = False
        
        cell_r2c1 = h_tbl.rows[2].cells[1]
        p0 = cell_r2c1.paragraphs[0]
        for r in list(p0.runs): p0._p.remove(r._r)
        r1 = p0.add_run("Objectif/Intitulé de la séquence pédagogique et volume horaire : ")
        r1.bold = True
        if len(cell_r2c1.paragraphs) > 1:
            p1 = cell_r2c1.paragraphs[1]
            for r in list(p1.runs): p1._p.remove(r._r)
        else:
            p1 = cell_r2c1.add_paragraph()
        r2 = p1.add_run(header_info['sequence_duree'])
        r2.bold = False
        
        # 2. Session Table
        start_s = page_idx * 4
        end_s = min(start_s + 4, total_sessions)
        page_sessions = sessions[start_s:end_s]
        
        for s_idx, session in enumerate(page_sessions):
            row = s_tbl.rows[s_idx + 1]
            set_cell_simple(row.cells[0], session['date_time'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[1], session['obj'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[2], session['desc'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[3], session['obs'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[4], session['disp'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[5], session['dur'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[6], session['cumul'], bold=True, sz_pt=8)
            set_cell_simple(row.cells[7], "", bold=False, sz_pt=8)
            
        # 3. Footer Table
        set_cell_simple(f_tbl.rows[0].cells[0], "", bold=False, sz_pt=8)
        set_cell_simple(f_tbl.rows[0].cells[1], f"Taux de réalisation du VH : {page_taux[page_idx]}", bold=True, sz_pt=8)
        
        set_cell_simple(f_tbl.rows[1].cells[0], "Commentaire du directeur sur le degré de réalisation des objectifs", bold=True, sz_pt=8)
        comm_text = f".................................................................... {page_comments[page_idx]}"
        set_cell_simple(f_tbl.rows[1].cells[1], comm_text, bold=True, sz_pt=8)
        
    doc.save(target_output)
    print(f"[OK] Document successfully built at: {target_output}")

if __name__ == '__main__':
    if len(sys.argv) < 4:
        print("Usage: build_module.py <master_template> <target_output> <json_path>")
        sys.exit(1)
    build_document(sys.argv[1], sys.argv[2], sys.argv[3])
