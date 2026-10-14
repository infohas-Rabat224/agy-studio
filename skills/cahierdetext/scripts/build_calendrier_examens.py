import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def build_calendrier_examens(target_dirs=None):
    """
    Génère le classeur Excel 'Calendrier des Examens' standardisé comprenant 3 feuilles :
    1. 1ère Année 2026-2027 (Synchronisé avec l'emploi du temps officiel d'accréditation)
    2. Format Image Corrigé (Gabarit conforme à l'image source avec millésimes 2026-2027 rectifiés)
    3. Modèle Référence (image.png) (Reproduction 100% fidèle du document visuel source)
    
    Sauvegarde automatiquement dans les répertoires cibles du projet Cahier de Texte / Planning.
    """
    wb = openpyxl.Workbook()

    # Typographie & Styles
    font_title = Font(name="Calibri", size=13, bold=True, color="000000")
    font_header = Font(name="Calibri", size=11, bold=True, color="000000")
    font_banner = Font(name="Calibri", size=11, bold=True, color="000000")
    font_module = Font(name="Calibri", size=10, bold=False, color="000000")
    font_nh = Font(name="Calibri", size=10, bold=True, color="000000")
    font_data = Font(name="Calibri", size=10, bold=False, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")

    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_total = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin_border_side = Side(style="thin", color="000000")
    double_border_side = Side(style="double", color="000000")
    
    border_box = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_border_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def apply_border_range(ws, start_row, end_row, start_col, end_col, border=border_box):
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border

    # -------------------------------------------------------------------------
    # Feuille 1: 1ère Année 2026-2027 (Synchronisé Emploi du Temps Officiel)
    # -------------------------------------------------------------------------
    ws_sync = wb.active
    ws_sync.title = "1ère Année 2026-2027"
    ws_sync.views.sheetView[0].showGridLines = True

    # -------------------------------------------------------------------------
    # Feuille 2: Format Image Corrigé (2026-2027)
    # -------------------------------------------------------------------------
    ws_corr = wb.create_sheet(title="Format Image Corrigé")
    ws_corr.views.sheetView[0].showGridLines = True

    # -------------------------------------------------------------------------
    # Feuille 3: Modèle Référence (image.png)
    # -------------------------------------------------------------------------
    ws_ref = wb.create_sheet(title="Modèle Référence (image.png)")
    ws_ref.views.sheetView[0].showGridLines = True

    # =========================================================================
    # JEUX DE DONNÉES PÉDAGOGIQUES ET EXAMENS
    # =========================================================================

    # 1. Modèle Référence (image.png brute)
    ref_modules = [
        (1, "Métier et Formation", 30, ["30-Sep-26"], "05-Apr-27", "29-Sep-26"),
        (5, "Comportement et Attitudes professionnels", 60, ["24-Nov-26", ""], "05-Apr-27", "23-Nov-26"),
        (7, "Approche Client", 60, ["24-Nov-26", ""], "6-avr.-27", "25-Nov-26"),
        (8, "Interaction en situations professionnels", 75, ["15-Oct-26", ""], "07-Apr-27", "16-Oct-26"),
        (3, "Communication en Français", 120, ["15-Dec-26"], "15-May-26", "16-Dec-26"),
        (4, "Environnement Professionnels", 180, ["01-Dec-26"], "14-May-26", "02-Dec-26"),
        (9, "Mesure de Sureté", 45, ["22-Dec-25"], "16-May-26", "23-Dec-25"),
        (10, "Marchandises Dangereuses", 45, ["05-Jan-26"], "16-May-26", "05-Jan-26"),
        (11, "Procédures de pré- départ", 60, ["19-janv.-26"], "16-May-26", "20-Jan-26"),
        (12, "Embarquement et Départ", 60, ["09-Feb-27"], "16-May-27", "10-Feb-27"),
        (2, "Outils Informatique", 60, ["2-mars-27"], "04-Mar-27", "03-Mar-27"),
        (6, "Communication en Anglais", 75, ["30-Mar-26"], "14-May-26", "31-Mar-27"),
        (13, "Activités d'accueil en milieu de travail(stage1)", 120, ["20/04/2027"], "16/05/2027", "20/04/2027")
    ]

    # 2. Format Image Corrigé (Millésimes 2026/2027 harmonisés)
    corr_modules = [
        (1, "Métier et Formation", 30, ["17-Jul-26"], "05-Apr-27", "17-Jul-26"),
        (5, "Comportement et Attitudes professionnels", 60, ["24-Nov-26", ""], "05-Apr-27", "23-Nov-26"),
        (7, "Approche Client", 60, ["24-Nov-26", ""], "06-Avr-27", "25-Nov-26"),
        (8, "Interaction en situations professionnels", 75, ["15-Oct-26", ""], "07-Apr-27", "16-Oct-26"),
        (3, "Communication en Français", 120, ["15-Dec-26"], "15-Mai-27", "16-Dec-26"),
        (4, "Environnement Professionnels", 180, ["01-Dec-26"], "14-Mai-27", "02-Dec-26"),
        (9, "Mesure de Sureté", 45, ["22-Dec-26"], "16-Mai-27", "23-Dec-26"),
        (10, "Marchandises Dangereuses", 45, ["05-Jan-27"], "16-Mai-27", "05-Jan-27"),
        (11, "Procédures de pré- départ", 60, ["19-Jan-27"], "16-Mai-27", "20-Jan-27"),
        (12, "Embarquement et Départ", 60, ["09-Feb-27"], "16-Mai-27", "10-Feb-27"),
        (2, "Outils Informatique", 60, ["02-Mar-27"], "04-Mar-27", "03-Mar-27"),
        (6, "Communication en Anglais", 75, ["30-Mar-27"], "14-Mai-27", "31-Mar-27"),
        (13, "Activités d'accueil en milieu de travail(stage1)", 120, ["20/04/2027"], "16/05/2027", "20/04/2027")
    ]

    # 3. Synchronisé Emploi du Temps Officiel (Emploi du temps 01 année accréditation 2026-2027)
    sync_modules = [
        (1, "Métier et Formation", 30, ["17-Jul-26"], "05-Apr-27", "17-Jul-26"),
        (5, "Comportement et Attitudes professionnels", 60, ["20-Jul-26", "23-Jul-26", "28-Jul-26"], "05-Apr-27", "01-Aou-26"),
        (7, "Approche Client", 60, ["02-Sep-26", "11-Sep-26"], "06-Avr-27", "18-Sep-26"),
        (8, "Interaction en situations professionnels", 75, ["05-Aou-26"], "07-Apr-27", "18-Sep-26"),
        (3, "Communication en Français", 120, ["29-Sep-26", "06-Oct-26", "12-Oct-26", "27-Oct-26", "01-Dec-26"], "15-Dec-26", "16-Dec-26"),
        (4, "Environnement Professionnels", 180, ["29-Sep-26", "12-Oct-26", "02-Nov-26", "09-Nov-26", "16-Nov-26", "21-Nov-26", "27-Nov-26"], "04-Dec-26", "04-Dec-26"),
        (9, "Mesure de Sureté", 45, ["02-Dec-26"], "16-Dec-26", "17-Dec-26"),
        (10, "Marchandises Dangereuses", 45, ["22-Dec-26", "28-Dec-26", "30-Dec-26"], "02-Jan-27", "08-Jan-27"),
        (11, "Procédures de pré- départ", 60, ["11-Jan-27", "13-Jan-27", "15-Jan-27"], "16-Jan-27", "22-Jan-27"),
        (12, "Embarquement et Départ", 60, ["05-Fev-27"], "12-Fev-27", "12-Fev-27"),
        (2, "Outils Informatique", 60, ["26-Fev-27"], "04-Mar-27", "05-Mar-27"),
        (6, "Communication en Anglais", 75, ["19-Mar-27"], "02-Avr-27", "02-Avr-27"),
        (13, "Activités d'accueil en milieu de travail(stage1)", 120, ["16/04/2027"], "23/04/2027", "23/04/2027")
    ]

    def populate_sheet(ws, modules, banner_text, include_total=True):
        # Ligne 1: Titre Principal
        ws.merge_cells("B1:E1")
        cell_t = ws.cell(row=1, column=2, value="CALENDRIER DES EXAMENS")
        cell_t.font = font_title
        cell_t.alignment = align_center

        # Ligne 2: En-têtes de colonnes
        headers = ["N°", "Titre de Module", "NH", "CC", "EFCF T", "EFCF P"]
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 24
        ws.row_dimensions[3].height = 20

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_box

        # Ligne 3: Bandeau Année
        ws.merge_cells("A3:F3")
        cell_b = ws.cell(row=3, column=1, value=banner_text)
        cell_b.font = font_banner
        cell_b.alignment = align_center
        cell_b.fill = fill_yellow
        apply_border_range(ws, 3, 3, 1, 6)

        current_row = 4
        for num, title, nh, cc_list, efcf_t, efcf_p in modules:
            num_rows = max(1, len(cc_list))
            start_r = current_row
            end_r = current_row + num_rows - 1

            # Remplissage des dates CC
            for idx in range(num_rows):
                r = start_r + idx
                ws.row_dimensions[r].height = 19
                cc_val = cc_list[idx] if idx < len(cc_list) else ""
                cell_cc = ws.cell(row=r, column=4, value=cc_val)
                cell_cc.font = font_data
                cell_cc.alignment = align_center

            # Colonnes principales
            cell_num = ws.cell(row=start_r, column=1, value=num)
            cell_num.font = font_data
            cell_num.alignment = align_center

            cell_title = ws.cell(row=start_r, column=2, value=title)
            cell_title.font = font_module
            cell_title.alignment = align_center

            cell_nh = ws.cell(row=start_r, column=3, value=nh)
            cell_nh.font = font_nh
            cell_nh.alignment = align_center

            cell_t = ws.cell(row=start_r, column=5, value=efcf_t)
            cell_t.font = font_data
            cell_t.alignment = align_center

            cell_p = ws.cell(row=start_r, column=6, value=efcf_p)
            cell_p.font = font_data
            cell_p.alignment = align_center

            # Fusion verticale si plusieurs dates de CC
            if num_rows > 1:
                ws.merge_cells(start_row=start_r, end_row=end_r, start_column=1, end_column=1)
                ws.merge_cells(start_row=start_r, end_row=end_r, start_column=2, end_column=2)
                ws.merge_cells(start_row=start_r, end_row=end_r, start_column=3, end_column=3)
                ws.merge_cells(start_row=start_r, end_row=end_r, start_column=5, end_column=5)
                ws.merge_cells(start_row=start_r, end_row=end_r, start_column=6, end_column=6)

            # Bordures complètes de la zone du module
            apply_border_range(ws, start_r, end_r, 1, 6)
            current_row = end_r + 1

        # Ligne de Totalisation (Volume Horaire 990h)
        if include_total:
            r_tot = current_row
            ws.row_dimensions[r_tot].height = 22
            ws.merge_cells(start_row=r_tot, end_row=r_tot, start_column=1, end_column=2)
            c_lbl = ws.cell(row=r_tot, column=1, value="TOTAL VOLUME HORAIRE")
            c_lbl.font = font_total
            c_lbl.alignment = align_center
            c_lbl.fill = fill_total

            c_sum = ws.cell(row=r_tot, column=3, value="=SUM(C4:C{})".format(r_tot - 1))
            c_sum.font = font_total
            c_sum.alignment = align_center
            c_sum.fill = fill_total

            for c_idx in [4, 5, 6]:
                c_empty = ws.cell(row=r_tot, column=c_idx, value="")
                c_empty.fill = fill_total

            apply_border_range(ws, r_tot, r_tot, 1, 6, border=border_total)

        # Largeurs de colonnes ajustées
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 44
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 16
        ws.column_dimensions['E'].width = 16
        ws.column_dimensions['F'].width = 16

    # Construction des 3 feuilles
    populate_sheet(ws_sync, sync_modules, "Premiére Année (2026/2027) - Synchronisé Emploi du Temps", include_total=True)
    populate_sheet(ws_corr, corr_modules, "Premiére Année (2026/2027) - Format Image Corrigé", include_total=True)
    populate_sheet(ws_ref, ref_modules, "Premiére Année", include_total=False)

    # Répertoires par défaut de sauvegarde
    if not target_dirs:
        target_dirs = [
            r"D:\Piopi\Cahier text\Planning année 2026-2028\1er année",
            r"D:\Piopi\Cahier text\Planning année 2026-2028"
        ]

    saved_files = []
    for d in target_dirs:
        if os.path.exists(d) or True:
            os.makedirs(d, exist_ok=True)
            for fname in ["CALENDRIER DES EXAMENS 1ère Année 2026-2027.xlsx", "Calendrier des Examens.xlsx"]:
                fpath = os.path.join(d, fname)
                wb.save(fpath)
                saved_files.append(fpath)
                print(f"[OK] Classeur Excel sauvegardé : {fpath}")

    return saved_files

if __name__ == "__main__":
    build_calendrier_examens()
