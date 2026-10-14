import docx
import os
import re
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn('w:shd'):
            tcPr.remove(child)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_properties(cell, text, font_name="Arial", font_size=8, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT, bg_hex=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 1.0
    
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
            
    if bg_hex:
        set_cell_background(cell, bg_hex)

def add_p_break(p):
    pPr = p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag == qn('w:pageBreakBefore'):
            pPr.remove(child)
    pb = OxmlElement('w:pageBreakBefore')
    pPr.append(pb)

def remove_all_breaks(elem):
    for pPr in elem.xpath('.//w:pPr'):
        for pb in list(pPr.xpath('./w:pageBreakBefore')):
            pPr.remove(pb)
    for r in elem.xpath('.//w:r'):
        for br in list(r.xpath('./w:br')):
            if br.get(qn('w:type')) == 'page':
                r.remove(br)

def build_teacher_classeur(config):
    t_key = config['key']
    t_name_upper = config['name_upper']
    t_name_title = config['name_title']
    t_specialite = config['specialite']
    out_docx = config['out_docx']
    modules_info = config['modules'] # list of 5 (num, title, hours, code)
    slots_data = config['weekly_slots']
    gantt1_hl = config['gantt1_highlight']
    gantt2_hl = config['gantt2_highlight']
    arch_data = config['architectures'] # list of 5 (title, seqs)
    eval_data = config['evaluations'] # list of 11

    src_base = r"C:\Users\piopi\.gemini\antigravity-cli\brain\1dcbb1ce-691f-4581-a6d0-c001a5605718\scratch\test_tight_end.docx"
    doc = docx.Document(src_base)
    body = doc._body._element

    # 1. Update Footers
    for rel_id, part in doc.part.related_parts.items():
        if 'footer' in part.partname or 'header' in part.partname:
            for t in part.element.xpath('.//*[local-name()="t"]'):
                if t.text:
                    if 'AZIZA JOUAIDI' in t.text:
                        t.text = t.text.replace('AZIZA JOUAIDI', t_name_upper)
                    if 'Aziza Jouaidi' in t.text:
                        t.text = t.text.replace('Aziza Jouaidi', t_name_title)

    # 2. Table 0 (Cover / Identification) - compact to strictly fit 1 page
    t0 = doc.tables[0]
    for r in t0.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    if run.font.size and run.font.size.pt > 9:
                        run.font.size = Pt(run.font.size.pt - 0.5)

    def set_t0_field(cell, text, font_size=8, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = bold

    # School details in Table 0 Row 4
    set_t0_field(t0.rows[4].cells[0], "Courriel : infohas@menara.ma  |  Site Web : www.infohas.ma", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Populate Formateur Identity & Units on Page 1 (Table 0)
    set_t0_field(t0.rows[11].cells[0], t_name_upper, font_size=9.5, bold=True)
    if 'address' in config:
        set_t0_field(t0.rows[12].cells[0], config['address'], font_size=8)
    if 'phone' in config:
        set_t0_field(t0.rows[13].cells[1], config['phone'], font_size=8)
    if 'email' in config:
        set_t0_field(t0.rows[14].cells[0], config['email'], font_size=8)
    if 'recrutement_date' in config:
        set_t0_field(t0.rows[16].cells[4], config['recrutement_date'], font_size=8)
    if 'statut' in config:
        set_t0_field(t0.rows[17].cells[4], config['statut'], font_size=8, bold=True)
    if 'diplome' in config:
        set_t0_field(t0.rows[19].cells[4], config['diplome'], font_size=8)
    if 'specialite' in config:
        set_t0_field(t0.rows[20].cells[3], config['specialite'], font_size=8)

    # Modules et volumes horaires assurés dans la filière
    if 't0_mod_1ere' in config:
        set_t0_field(t0.rows[27].cells[5], config['t0_mod_1ere'], font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if 't0_mod_2eme' in config:
        set_t0_field(t0.rows[27].cells[12], config['t0_mod_2eme'], font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if 't0_vol_1ere' in config:
        set_t0_field(t0.rows[28].cells[5], config['t0_vol_1ere'], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    if 't0_vol_2eme' in config:
        set_t0_field(t0.rows[28].cells[12], config['t0_vol_2eme'], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Années de formation (Cycle 2025 - 2027)
    set_t0_field(t0.rows[29].cells[5], "2025 / 2026", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_t0_field(t0.rows[29].cells[12], "2026 / 2027", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Fallback cleanup across all Table 0 cells for unique legacy strings
    repl_map = {
        "INFOHAS@MENARA.MA": "Courriel : infohas@menara.ma  |  Site Web : www.infohas.ma"
    }
    if 'address' in config:
        repl_map["Hay Al Maghrib AlArabi n°2915 Massira II Temara"] = config['address']
    if 'phone' in config:
        repl_map["0537661763"] = config['phone']
    if 'email' in config:
        repl_map["Ajouaidi999@gmail.com"] = config['email']

    for r in t0.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    for old_k, new_v in repl_map.items():
                        if old_k in run.text and run.text != new_v:
                            run.text = run.text.replace(old_k, new_v)


    # 3. Table 1 (Horaires formateur)
    t1 = doc.tables[1]
    set_cell_properties(t1.rows[0].cells[1], f"HORAIRES D'INTERVENTION DU FORMATEUR ({t_name_upper})", font_name="Arial", font_size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_properties(t1.rows[1].cells[1], "09h00 – 12h00 (Matinée)", font_name="Arial", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_properties(t1.rows[1].cells[2], "13h00 – 16h00 (Après-midi)", font_name="Arial", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_properties(t1.rows[1].cells[3], "Samedi (09h00 – 12h00)", font_name="Arial", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r_idx, (j, m, a, s) in enumerate(slots_data, start=2):
        row = t1.rows[r_idx]
        set_cell_properties(row.cells[0], j, font_name="Arial", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_properties(row.cells[1], m, font_name="Arial", font_size=7.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_properties(row.cells[2], a, font_name="Arial", font_size=7.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_properties(row.cells[3], s, font_name="Arial", font_size=7.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 4. Table 3 & 4 (Gantt Charts)
    t3 = doc.tables[3]
    t3.rows[1].cells[0].text = "1ère  année de formation (2025 / 2026) — Filière ATAM"
    t3.rows[1].cells[0].paragraphs[0].runs[0].font.name = "Arial"
    t3.rows[1].cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
    t3.rows[1].cells[0].paragraphs[0].runs[0].font.bold = True
    for r in t3.rows[2:]:
        lbl = r.cells[0].text
        is_mine = any(k in lbl for k in gantt1_hl)
        lbl_bg = "B4C6E7" if is_mine else "F9F9F9"
        set_cell_background(r.cells[0], lbl_bg)
        for c_idx in range(1, 11):
            txt = r.cells[c_idx].text.strip()
            if "■" in txt:
                fill = "8EA9DB" if is_mine else "D9D9D9"
                set_cell_properties(r.cells[c_idx], "■", font_name="Arial", font_size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex=fill)

    # Paragraph before Table 4
    for p in doc.paragraphs:
        if "Planning de réalisation des unités de formation — 2ème Année de formation" in p.text:
            p.text = "Planning de réalisation des unités de formation — 2ème Année de formation (2026 / 2027)"
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    t4 = doc.tables[4]
    t4.rows[1].cells[0].text = "2ème  année de formation (2026 / 2027) — Filière ATAM"
    t4.rows[1].cells[0].paragraphs[0].runs[0].font.name = "Arial"
    t4.rows[1].cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
    t4.rows[1].cells[0].paragraphs[0].runs[0].font.bold = True
    for r in t4.rows[2:]:
        lbl = r.cells[0].text
        is_mine = any(k in lbl for k in gantt2_hl)
        lbl_bg = "B4C6E7" if is_mine else "F9F9F9"
        set_cell_background(r.cells[0], lbl_bg)
        for c_idx in range(1, 11):
            txt = r.cells[c_idx].text.strip()
            if "■" in txt:
                fill = "8EA9DB" if is_mine else "D9D9D9"
                set_cell_properties(r.cells[c_idx], "■", font_name="Arial", font_size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex=fill)

    # 5. Table 5 (Emploi du temps type filière) - compact so NB stays on Page 10
    t5 = doc.tables[5]
    t5.rows[0].cells[0].text = "EMPLOI DU TEMPS HEBDOMADAIRE TYPE DE LA FILIÈRE ATAM - 1ère ANNÉE (2025 / 2026)"
    t5.rows[0].cells[0].paragraphs[0].runs[0].font.name = "Arial"
    t5.rows[0].cells[0].paragraphs[0].runs[0].font.size = Pt(8)
    t5.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True
    for r in t5.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0.5)
                p.paragraph_format.space_after = Pt(0.5)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(7.5)
    t5.rows[7].cells[1].text = "09h00 – 12h00 : Contrôles Continus (C.C) & Épreuves de Fin de Module (EFCF)\nM5 Comportements (02/08/25) | M11 Pré-départ (23/01/26) | M4 Environnement (Dec 25)"
    t5.rows[7].cells[1].paragraphs[0].runs[0].font.size = Pt(7.5)

    # Find NB after Table 5 and ensure no page breaks
    for p in doc.paragraphs:
        if "NB : d’autres informations peuvent y figurer" in p.text:
            remove_all_breaks(p._p)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                run.font.size = Pt(7.5)

    # 6. Table 6 (Stagiaires)
    t6 = doc.tables[6]
    t6.rows[1].cells[2].text = "2025 / 2026"

    # 7. Dividers (P13-P17)
    div_pairs = [(169, 171), (196, 198), (216, 218), (236, 238), (244, 246)]
    for (p_num_idx, p_title_idx), (m_num, m_title, m_h, m_code) in zip(div_pairs, modules_info):
        p_num = doc.paragraphs[p_num_idx]
        p_title = doc.paragraphs[p_title_idx]
        p_num.text = f"UNITE DE FORMATION N° {m_num}"
        p_num.runs[0].font.name = "Arial"
        p_num.runs[0].font.size = Pt(14)
        p_num.runs[0].font.bold = True
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p_title.text = f"« {m_title} » ({m_h}h)"
        p_title.runs[0].font.name = "Arial"
        p_title.runs[0].font.size = Pt(12)
        p_title.runs[0].font.bold = True
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 8. Architecture Tables (Tables 7-11)
    for t_idx, (tbl_title, seqs) in zip(range(7, 12), arch_data):
        tbl = doc.tables[t_idx]
        set_cell_properties(tbl.rows[0].cells[0], tbl_title, font_name="Arial", font_size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex="D9E1F2")
        col_count = len(tbl.rows[1].cells)
        for c_idx in range(col_count):
            seq_text = seqs[c_idx] if c_idx < len(seqs) else ""
            set_cell_properties(tbl.rows[1].cells[c_idx], seq_text, font_name="Arial", font_size=7, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 9. Clean empty paragraphs between Architecture Tables (Tables 7-11)
    # and ensure clean PageBreakBefore on each title paragraph
    for i in range(7, 11):
        t_curr = doc.tables[i]._element
        t_next = doc.tables[i+1]._element
        curr = t_curr.getnext()
        while curr is not None and curr != t_next:
            nxt = curr.getnext()
            if curr.tag.split('}')[-1] == 'p':
                txt = ''.join(curr.itertext()).strip()
                if not txt: # empty paragraph
                    body.remove(curr)
                elif "architecture de" in txt.lower():
                    add_p_break(curr)
            curr = nxt

    # Between Table 11 and Table 12: keep exactly ONE paragraph with pageBreakBefore
    t11_elem = doc.tables[11]._element
    t12_elem = doc.tables[12]._element
    curr = t11_elem.getnext()
    p_before_t12 = None
    while curr is not None and curr != t12_elem:
        nxt = curr.getnext()
        if curr.tag.split('}')[-1] == 'p':
            if p_before_t12 is None:
                p_before_t12 = curr
                add_p_break(p_before_t12)
                p_before_t12.get_or_add_pPr()
            else:
                body.remove(curr)
        curr = nxt

    # After Table 12: remove all empty paragraphs before the sectPr paragraph
    curr = t12_elem.getnext()
    while curr is not None and curr.tag.split('}')[-1] == 'p':
        nxt = curr.getnext()
        txt = ''.join(curr.itertext()).strip()
        has_sect = len(curr.xpath('.//w:sectPr')) > 0
        if not txt and not has_sect:
            body.remove(curr)
        elif has_sect or txt:
            break
        curr = nxt

    # Ensure cantSplit on all rows of Tables 7-12
    for i in range(7, 13):
        for r in doc.tables[i].rows:
            rPr = r._tr.get_or_add_trPr()
            if not rPr.xpath('./w:cantSplit'):
                rPr.append(OxmlElement('w:cantSplit'))

    # 10. Evaluation Sheets (Tables 13 to 23)
    for idx, ed in enumerate(eval_data, start=13):
        t = doc.tables[idx]
        for r in t.rows:
            rPr = r._tr.get_or_add_trPr()
            if not rPr.xpath('./w:cantSplit'):
                rPr.append(OxmlElement('w:cantSplit'))
            h_nodes = rPr.xpath('./w:trHeight')
            for hn in h_nodes:
                if int(hn.get(qn('w:val'), 0)) > 600:
                    hn.set(qn('w:val'), '550')

        t_prof, t_mod, t_seq, t_obj, t_date, t_dur, t_ep, t_cond, t_crit, t_desc = ed
        set_cell_properties(t.rows[2].cells[0], f"Formateur : {t_prof}", font_name="Arial", font_size=8, bold=True)
        set_cell_properties(t.rows[2].cells[2], f"Unité de formation : {t_mod}", font_name="Arial", font_size=8, bold=True)
        set_cell_properties(t.rows[3].cells[0], f"Séquence pédagogique : {t_seq}", font_name="Arial", font_size=7.5)
        set_cell_properties(t.rows[3].cells[2], f"Objectif pédagogique : {t_obj}", font_name="Arial", font_size=7.5)
        set_cell_properties(t.rows[5].cells[0], f"Date : {t_date}", font_name="Arial", font_size=8, bold=True)
        set_cell_properties(t.rows[5].cells[1], f"Durée : {t_dur}", font_name="Arial", font_size=8)
        set_cell_properties(t.rows[5].cells[2], f"Nature et numéro de l'épreuve : {t_ep}", font_name="Arial", font_size=8, bold=True)
        set_cell_properties(t.rows[7].cells[0], f"Conditions et organisation du contrôle : {t_cond}", font_name="Arial", font_size=7)
        set_cell_properties(t.rows[7].cells[3], f"Critères d’appréciation des résultats : {t_crit}", font_name="Arial", font_size=7)
        set_cell_properties(t.rows[8].cells[0], f"Description des objectifs de l’épreuve : {t_desc}", font_name="Arial", font_size=7)

        # Explicit page break before each evaluation sheet (T14 to T23)
        if idx > 13:
            prev = t._element.getprevious()
            if prev is not None and prev.tag.split('}')[-1] == 'p':
                add_p_break(prev)

    # 11. Table 24 (Synthèse des résultats) - compact to strictly fit Page 36
    t24 = doc.tables[24]
    set_cell_properties(t24.rows[2].cells[0], f"Formateur : {t_name_title}", font_name="Arial", font_size=8, bold=True)
    mod_summary = ", ".join(m[3] for m in modules_info)
    set_cell_properties(t24.rows[2].cells[2], f"Unité de formation : Modules {mod_summary}", font_name="Arial", font_size=8, bold=True)
    set_cell_properties(t24.rows[5].cells[0], "Date : Session 2025 / 2026 - 2026 / 2027", font_name="Arial", font_size=8, bold=True)
    
    # Compact student rows of Table 24 (rows 8 to 33)
    for r_idx in range(8, 34):
        r = t24.rows[r_idx]
        trPr = r._tr.get_or_add_trPr()
        for child in list(trPr):
            if child.tag == qn('w:trHeight'): trPr.remove(child)
        trH = OxmlElement('w:trHeight')
        trH.set(qn('w:val'), '240') # 12pt
        trH.set(qn('w:hRule'), 'atLeast')
        trPr.append(trH)
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(9)
                if not p.runs:
                    run = p.add_run()
                else: run = p.runs[0]
                run.font.size = Pt(6.5)

    for r in t24.rows:
        rPr = r._tr.get_or_add_trPr()
        if not rPr.xpath('./w:cantSplit'):
            rPr.append(OxmlElement('w:cantSplit'))

    # Ensure FICHE SYNTHETIQUE has exactly ONE page break before it
    # Note: In the base docx, the heading run already has <w:br w:type="page"/>
    # Adding pageBreakBefore causes a DOUBLE page break pushing it to Page 37!
    # So we remove any duplicate pageBreakBefore on p_curr and intermediate paragraphs.
    t24_elem = t24._element
    p_curr = t24_elem.getprevious()
    while p_curr is not None:
        txt = ''.join(p_curr.itertext()).strip()
        if "FICHE SYNTHETIQUE" in txt:
            # Remove any pageBreakBefore so it doesn't double break with the inline <w:br>
            pPr = p_curr.get_or_add_pPr()
            for pb in list(pPr.xpath('./w:pageBreakBefore')):
                pPr.remove(pb)
            has_br = len(p_curr.xpath('.//w:br[@w:type="page"]')) > 0
            if not has_br:
                add_p_break(p_curr)
            break
        else:
            # remove page break on intermediate paragraphs
            remove_all_breaks(p_curr)
        p_curr = p_curr.getprevious()

    # 12. Table 28 & 29 (Stages)
    for r in doc.tables[28].rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    if "2026 / 2027" in run.text: run.text = run.text.replace("2026 / 2027", "2025 / 2026")
                    if "2027 / 2028" in run.text: run.text = run.text.replace("2027 / 2028", "2026 / 2027")

    for r in doc.tables[29].rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    if "29/03/2027 au 16/04/2027" in run.text: run.text = run.text.replace("29/03/2027 au 16/04/2027", "30/03/2026 au 17/04/2026")
                    if "2027 / 2028" in run.text: run.text = run.text.replace("2027 / 2028", "2026 / 2027")

    # Tighten last paragraph
    last_p = doc.paragraphs[-1]
    last_p.paragraph_format.space_before = Pt(0)
    last_p.paragraph_format.space_after = Pt(0)
    last_p.paragraph_format.line_spacing = Pt(1)
    if not last_p.runs:
        r = last_p.add_run()
    else:
        r = last_p.runs[0]
    r.font.size = Pt(1)

    doc.save(out_docx)
    print(f"[{t_key}] Document saved to {out_docx}")

