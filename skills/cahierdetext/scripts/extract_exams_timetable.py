import os
import sys
import docx

def extract_exams_from_timetable(docx_path):
    """
    Scans an official timetable docx (e.g. 'Emploi du temps 01 année acreditation 2026-2027.docx')
    and extracts all evaluation sessions (Contrôles Continus CC, EFCF Théorique, EFCF Pratique, Examens).
    """
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Fichier emploi du temps introuvable : {docx_path}")

    doc = docx.Document(docx_path)
    print(f"Extraction des examens depuis : {docx_path}")
    print(f"Nombre de tables détectées : {len(doc.tables)}")

    all_sessions = []
    for t_idx, table in enumerate(doc.tables):
        week_str = table.rows[0].cells[1].text.strip() if len(table.rows) > 0 and len(table.rows[0].cells) > 1 else f"Semaine {t_idx+1}"
        for r_idx in range(2, len(table.rows)):
            c = table.rows[r_idx].cells
            jour = c[0].text.strip().replace('\n', ' ')
            horaire = c[1].text.strip().replace('\n', ' ') if len(c) > 1 else ''
            mod = c[2].text.strip().replace('\n', ' ') if len(c) > 2 else ''
            obj = c[3].text.strip().replace('\n', ' ') if len(c) > 3 else ''
            prof = c[4].text.strip().replace('\n', ' ') if len(c) > 4 else ''
            all_sessions.append({'week': week_str, 'jour': jour, 'horaire': horaire, 'mod': mod, 'obj': obj, 'prof': prof, 'table': t_idx+1})

    # Regroupement par module
    modules = {}
    for s in all_sessions:
        m = s['mod']
        if m not in modules:
            modules[m] = []
        modules[m].append(s)

    eval_data = {}
    for m, sess in modules.items():
        evals = []
        for s in sess:
            is_eval = any(k in s['obj'].upper() for k in ['C.C', 'CC', 'EFCF', 'EVAL', 'EXAM']) or any(k in s['mod'].lower() for k in ['examen', 'soutenance'])
            if is_eval:
                evals.append(s)
        if evals:
            eval_data[m] = evals

    return eval_data

if __name__ == "__main__":
    default_path = r"D:\Piopi\Cahier text\Planning année 2026-2028\1er année\Emploi du temps 01 année acreditation 2026-2027.docx"
    path = sys.argv[1] if len(sys.argv) > 1 else default_path
    res = extract_exams_from_timetable(path)
    for mod, evs in res.items():
        print(f"\nMODULE: {mod}")
        for ev in evs:
            print(f"  --> {ev['week']} | {ev['jour']} | {ev['horaire']} | {ev['obj']} | {ev['prof']}")
