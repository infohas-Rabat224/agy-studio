"""
Automated QA Verification and PDF Export script for Cahier de Texte documents.
Run with:
  uv run --with python-docx --with pywin32 --with pypdf python verify_module.py <docx_path> <expected_json> <expected_pages> [pdf_path]
"""
import sys
import os
import json
import docx
import win32com.client
import pypdf

def verify_module(docx_path, json_path, expected_pages, pdf_path=None):
    docx_path = os.path.abspath(docx_path)
    json_path = os.path.abspath(json_path)
    expected_pages = int(expected_pages)
    
    if pdf_path is None:
        pdf_path = os.path.splitext(docx_path)[0] + '.pdf'
    else:
        pdf_path = os.path.abspath(pdf_path)

    print(f"--- Verifying: {docx_path} ---")
    doc = docx.Document(docx_path)
    
    expected_tables = expected_pages * 3
    if len(doc.tables) != expected_tables:
        print(f"[ERROR] Table count mismatch: expected {expected_tables}, got {len(doc.tables)}")
        return False
    print(f"[OK] Table count matches: {expected_tables}")
    
    # 1. Timetable dates & cumul
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    expected_sessions = data['sessions']
    
    found_sessions = []
    for p_idx in range(expected_pages):
        s_tbl = doc.tables[p_idx * 3 + 1]
        for r_idx in range(1, len(s_tbl.rows)):
            c = s_tbl.rows[r_idx].cells
            found_sessions.append({
                'date': c[0].text.strip(),
                'obj': c[1].text.strip(),
                'dur': c[5].text.strip(),
                'cumul': c[6].text.strip()
            })
            
    if len(found_sessions) != len(expected_sessions):
        print(f"[ERROR] Session count mismatch: expected {len(expected_sessions)}, got {len(found_sessions)}")
        return False
        
    mismatches = 0
    for i in range(len(expected_sessions)):
        exp = expected_sessions[i]
        fnd = found_sessions[i]
        if exp['date_time'] != fnd['date']:
            print(f"[MISMATCH] S{i+1} Date: expected '{exp['date_time']}', got '{fnd['date']}'")
            mismatches += 1
        if exp['cumul'] != fnd['cumul']:
            print(f"[MISMATCH] S{i+1} Cumul: expected '{exp['cumul']}', got '{fnd['cumul']}'")
            mismatches += 1
            
    if mismatches == 0:
        print(f"[OK] All {len(expected_sessions)} sessions match dates and cumulative hours exactly.")
    else:
        return False

    # 2. Word COM Page Count & PDF Export
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(docx_path)
        word_pages = wdoc.ComputeStatistics(2) # 2 = wdStatisticPages
        print(f"[Word COM] Detected Pages: {word_pages} (Expected: {expected_pages})")
        wdoc.SaveAs2(pdf_path, 17) # 17 = wdFormatPDF
        wdoc.Close(False)
    finally:
        try:
            word.Quit()
        except Exception:
            pass
        
    if word_pages != expected_pages:
        print(f"[ERROR] Word page count mismatch! Expected {expected_pages}, got {word_pages}")
        return False
        
    # 3. PDF Page Count
    reader = pypdf.PdfReader(pdf_path)
    pdf_pages = len(reader.pages)
    print(f"[PDF Verification] Pages in PDF: {pdf_pages}")
    if pdf_pages != expected_pages:
        print(f"[ERROR] PDF page count mismatch! Expected {expected_pages}, got {pdf_pages}")
        return False

    print(f"[SUCCESS] Document {docx_path} verified 100% and exported to {pdf_path}")
    return True

if __name__ == '__main__':
    if len(sys.argv) < 4:
        print("Usage: verify_module.py <docx_path> <expected_json> <expected_pages> [pdf_path]")
        sys.exit(1)
    pdf = sys.argv[4] if len(sys.argv) > 4 else None
    success = verify_module(sys.argv[1], sys.argv[2], sys.argv[3], pdf)
    sys.exit(0 if success else 1)
